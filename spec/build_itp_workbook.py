"""Builds InvoiceToPay_BuildWorkbook.xlsx from itp_data.py (the single source)."""
import re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import ColorScaleRule
import itp_data as D

INK, HEAD, BAND, EDIT, ACC = "1F2937", "1F2937", "F3F4F6", "FFF2CC", "B7791F"
F_TITLE = Font(name="Arial", size=14, bold=True, color=INK)
F_SUB = Font(name="Arial", size=10, italic=True, color="6B7280")
F_HEAD = Font(name="Arial", size=10, bold=True, color="FFFFFF")
F_BODY = Font(name="Arial", size=10, color=INK)
F_ID = Font(name="Arial", size=10, bold=True, color=ACC)
F_CODE = Font(name="Courier New", size=9, color=INK)
FILL_HEAD = PatternFill("solid", fgColor=HEAD)
FILL_BAND = PatternFill("solid", fgColor=BAND)
FILL_EDIT = PatternFill("solid", fgColor=EDIT)
THIN = Side(style="thin", color="D1D5DB")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
HR = 4          # header row
R0 = 5          # first data row
LAST = 1000     # formula range end

wb = Workbook()
wb.remove(wb.active)

def sheet(name, title, subtitle, headers, widths, rows, id_cols=(1,), code_cols=(), edit_cols=(), freeze_col=2):
    ws = wb.create_sheet(name)
    ws["A1"] = title; ws["A1"].font = F_TITLE
    ws["A2"] = subtitle; ws["A2"].font = F_SUB
    for c, (h, w) in enumerate(zip(headers, widths), 1):
        cell = ws.cell(HR, c, h); cell.font = F_HEAD; cell.fill = FILL_HEAD
        cell.alignment = Alignment(wrap_text=True, vertical="center"); cell.border = BORDER
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.row_dimensions[HR].height = 30
    for r, row in enumerate(rows, R0):
        for c, v in enumerate(row, 1):
            cell = ws.cell(r, c, v)
            cell.font = F_CODE if c in code_cols else (F_ID if c in id_cols else F_BODY)
            cell.alignment = WRAP; cell.border = BORDER
            if c in edit_cols: cell.fill = FILL_EDIT
            elif (r - R0) % 2: cell.fill = FILL_BAND
    ws.freeze_panes = ws.cell(R0, freeze_col)
    if rows:
        ws.auto_filter.ref = f"A{HR}:{get_column_letter(len(headers))}{R0 + len(rows) - 1}"
    return ws

# ---------------- derived lookups (single source) ----------------
task_of = {t[0]: t for t in D.TASKS}
cs_by_task, tests_by_task, tests_by_cs = {}, {}, {}
for cs in D.CODE: cs_by_task.setdefault(cs[1], []).append(cs[0])
for t in D.TESTS:
    tests_by_task.setdefault(t[4], []).append(t[0]); tests_by_cs.setdefault(t[3], []).append(t[0])
split = lambda s: [x.strip() for x in s.split(",") if x.strip()]
contracts_by_task, cs_by_model = {}, {}
for cs in D.CODE:
    for m in split(cs[9]) + split(cs[10]):
        contracts_by_task.setdefault(cs[1], set()).add(m); cs_by_model.setdefault(m, []).append(cs[0])
model_order = {m[1]: i for i, m in enumerate(D.MODELS)}

def class_src(m):
    _id, name, module, kind, purpose, base, config, validators, task = m
    gen = re.match(r"(\w+)\[(\w+)\]", name)
    cname = gen.group(1) if gen else name
    bases = base if not gen else base
    lines = [f"# {module}", f"class {cname}({bases}):", f'    """{purpose}."""']
    if config and config != "—":
        cfg = "SettingsConfigDict" if "BaseSettings" in base else "ConfigDict"
        lines.append(f"    model_config = {cfg}({config})")
    for f, typ, default, cons, desc in D.FIELDS[name]:
        line = f"    {f}: {typ}"
        if default != "required": line += f" = {default}"
        note = "; ".join(x for x in (cons, desc) if x)
        if note: line += f"  # {note}"
        lines.append(line)
    if validators and validators != "—":
        lines.append(f"    # invariant: {validators}")
    return "\n".join(lines)

# ---------------- Index ----------------
ws = wb.create_sheet("Index")
ws["A1"] = D.META["title"]; ws["A1"].font = F_TITLE
ws["A2"] = D.META["subtitle"]; ws["A2"].font = F_SUB
ws["A3"] = f"Generated from {D.META['source']} · package {D.META['package']} · edit the data module and rebuild, never the cells (except the yellow ones)."; ws["A3"].font = F_SUB
SHEETS = [
    ("Stack", "Every runtime, framework and tool with package, minimum line and a column for the resolved version"),
    ("Framework Rules", "Per framework: where it may be imported, what to use it for and not, APIs, what you implement, gotchas"),
    ("Conventions", "Coding conventions Claude Code must follow"),
    ("Guardrails", "What Claude Code must never do, and why"),
    ("Definition of Done", "The checklist every task must satisfy"),
    ("Local Services", "docker compose services and processes for local development"),
    ("Env Variables", "Every ITP_ environment variable, derived from the Settings contract"),
    ("Commands", "Make targets and the commands behind them"),
    ("Bootstrap Files", "Files created at the start and what goes in them"),
    ("Prompt Registry", "Every prompt id, the code that uses it, its model role and output schema"),
    ("Mail Templates", "Approved supplier mail templates"),
    ("Fixtures", "Seed and test data (fictitious) for the stub ERP and scenarios"),
    ("Open Decisions", "What must be decided, with the default Claude Code uses until then"),
    ("Glossary", "Terms used across the workbook"),
    ("Task Prompts", "A ready-to-paste Claude Code prompt per task, generated from all the sheets"),
    ("Repo Docs", "The Markdown files pushed with the code so any session can continue: generated vs living, read when, updated when"),
    ("Session Groups", "16 groups of tasks, one branch each, with the model to use and live days"),
    ("Session Protocol", "What every Claude Code session does at start, during and at end, and the slash commands that do it"),
    ("Increments", "Six increments with live task counts, estimated days and progress"),
    ("Tasks", "Every build task: objective, framework features, controls, derived code sections, contracts and tests; status to fill in"),
    ("Code Sections", "Every class, function, route, node, tool and test module, with signature, framework feature and contracts in/out"),
    ("Modules", "The package layout: one row per module with its layer and section count"),
    ("Contracts", "Every Pydantic model with its kind, config, invariants and a generated class skeleton"),
    ("Contract Fields", "Every field of every model: type, default, constraints, meaning"),
    ("Ports & Adapters", "Each port, its methods, the adapter and the framework features behind it"),
    ("Tool Registry", "Every tool with risk class, models, preconditions, idempotency key, verification and undo"),
    ("API Routes", "Every HTTP route with request/response models, dependencies and status codes"),
    ("Run Graph", "Nodes and edges of the run, resolution and answer graphs"),
    ("Controls", "The 15 controls mapped to the code sections that implement them"),
    ("Recovery Policy", "Typed error → response, limit, target"),
    ("Autonomy Matrix", "What runs alone, what needs approval, what never runs"),
    ("Persistence", "Tables and stores, their owner, keys and retention"),
    ("Events", "Domain events, producers, consumers, transport"),
    ("Tests", "Every test with its target code section and task"),
    ("Traceability", "Control capabilities × program capabilities, counted live from Tasks"),
]
for c, (h, w) in enumerate([("Sheet", 22), ("What it holds", 90), ("Rows", 10)], 1):
    cell = ws.cell(5, c, h); cell.font = F_HEAD; cell.fill = FILL_HEAD; cell.border = BORDER
    ws.column_dimensions[get_column_letter(c)].width = w
for r, (s, what) in enumerate(SHEETS, 6):
    ws.cell(r, 1, s).font = F_ID; ws.cell(r, 2, what).font = F_BODY
    if s not in ("Traceability",):
        ws.cell(r, 3, f"=COUNTA('{s}'!A{R0}:A{LAST})").font = F_BODY
    else:
        ws.cell(r, 3, len(D.CONTROL_CAPS)).font = F_BODY
    for c in (1, 2, 3): ws.cell(r, c).border = BORDER; ws.cell(r, c).alignment = WRAP
r = 6 + len(SHEETS) + 1
ws.cell(r, 1, "Progress").font = Font(name="Arial", size=11, bold=True, color=INK)
stats = [
    ("Tasks", f"=COUNTA(Tasks!A{R0}:A{LAST})", None),
    ("Estimated days", f"=SUM(Tasks!M{R0}:M{LAST})", "0.0"),
    ("Tasks done", f"=COUNTIF(Tasks!O{R0}:O{LAST},\"Done\")", None),
    ("Progress", f"=IF(B{r+1}=0,0,B{r+3}/B{r+1})", "0.0%"),
]
for i, (k, f, fmt) in enumerate(stats, 1):
    ws.cell(r + i, 1, k).font = F_BODY
    c = ws.cell(r + i, 2, f); c.font = F_BODY
    if fmt: c.number_format = fmt
r += len(stats) + 2
ws.cell(r, 1, "Legend").font = Font(name="Arial", size=11, bold=True, color=INK)
legend = [
    ("Yellow cells", "The only cells to edit: Stack resolved versions, Open Decisions, and the optional Status/Notes mirror on Tasks. The real build status is docs/PROGRESS.md in the repo."),
    ("Gold IDs", "T- task · CS- code section · M- contract model · TS- test · CAP- capability · INC- increment"),
    ("Derived columns", "Code sections, contracts and tests per task are derived from the data module, never typed twice"),
    ("Example", "T-206 Three-way match → CS-045 three_way_match() and CS-046 match node → contracts Invoice, PurchaseOrder, MatchResult … → tests TS-08, TS-09"),
]
for i, (k, v) in enumerate(legend, 1):
    ws.cell(r + i, 1, k).font = F_ID; c = ws.cell(r + i, 2, v); c.font = F_BODY; c.alignment = WRAP
ws.cell(r + 1, 1).fill = FILL_EDIT
r += len(legend) + 2
ws.cell(r, 1, "Start with Claude Code").font = Font(name="Arial", size=11, bold=True, color=INK)
steps = [
    ("1", "Record decisions on Open Decisions, or accept the defaults (at least OD-03 model ids before T-111)."),
    ("2", "Unzip invoice_to_pay_repo_starter.zip as the repo root (CLAUDE.md, docs/, tasks/, spec/, scripts/, .claude/) and commit it."),
    ("3", "In Claude Code: /start-session, then /next-group (switch /model to the group's model). End every session with /end-session."),
    ("4", "Build status lives in docs/PROGRESS.md (pushed with the code). The Status column here is optional."),
    ("5", "Spec change: edit spec/itp_data.py, log it in docs/SPEC_CHANGES.md, run make docs (living docs are never overwritten), then change code."),
]
for i, (k, v) in enumerate(steps, 1):
    ws.cell(r + i, 1, k).font = F_ID; c = ws.cell(r + i, 2, v); c.font = F_BODY; c.alignment = WRAP

# ---------------- Increments ----------------
inc_rows = [(i, n, g, d) for i, n, g, d in D.INCREMENTS]
ws = sheet("Increments", "Increments", "Each increment ends with a working program; counts and progress update from the Tasks sheet.",
           ["Increment", "Name", "Goal", "Done when", "Tasks", "Est. days", "Done", "Progress"],
           [11, 30, 50, 50, 9, 10, 9, 11], inc_rows)
for r in range(R0, R0 + len(inc_rows)):
    ws.cell(r, 5, f"=COUNTIF(Tasks!$B${R0}:$B${LAST},A{r})")
    ws.cell(r, 6, f"=SUMIF(Tasks!$B${R0}:$B${LAST},A{r},Tasks!$M${R0}:$M${LAST})").number_format = "0.0"
    ws.cell(r, 7, f"=COUNTIFS(Tasks!$B${R0}:$B${LAST},A{r},Tasks!$O${R0}:$O${LAST},\"Done\")")
    ws.cell(r, 8, f"=IF(E{r}=0,0,G{r}/E{r})").number_format = "0.0%"
    for c in range(5, 9): ws.cell(r, c).font = F_BODY; ws.cell(r, c).border = BORDER
tr = R0 + len(inc_rows)
ws.cell(tr, 2, "Total").font = Font(name="Arial", size=10, bold=True)
for c, f in ((5, f"=SUM(E{R0}:E{tr-1})"), (6, f"=SUM(F{R0}:F{tr-1})"), (7, f"=SUM(G{R0}:G{tr-1})"), (8, f"=IF(E{tr}=0,0,G{tr}/E{tr})")):
    cell = ws.cell(tr, c, f); cell.font = Font(name="Arial", size=10, bold=True)
    cell.number_format = "0.0%" if c == 8 else ("0.0" if c == 6 else "General")

# ---------------- Tasks ----------------
cap_name = {c[0]: c[1] for c in D.CAPABILITIES}
task_rows = []
for t in D.TASKS:
    tid, inc, cap, title, obj, fw, states, ccaps, dep, days, done = t
    task_rows.append((tid, inc, f"{cap} {cap_name[cap]}", title, obj, fw, states, ccaps,
                      ", ".join(cs_by_task.get(tid, [])),
                      ", ".join(sorted(contracts_by_task.get(tid, []), key=lambda m: model_order.get(m, 999))),
                      ", ".join(tests_by_task.get(tid, [])), dep, days, done, "Not started", ""))
ws = sheet("Tasks", "Tasks", "One row per build task. Code sections, contracts and tests are derived from the code-section and test data.",
           ["ID", "Increment", "Capability", "Title", "Objective", "Frameworks and features", "Control states", "Control capabilities",
            "Code sections", "Contracts touched", "Tests", "Depends on", "Est. days", "Done when", "Status", "Notes"],
           [8, 9, 24, 28, 48, 38, 18, 30, 22, 36, 14, 12, 8, 40, 13, 24], task_rows, edit_cols=(15, 16))
dv = DataValidation(type="list", formula1='"Not started,In progress,Blocked,Done"', allow_blank=True)
ws.add_data_validation(dv); dv.add(f"O{R0}:O{R0 + len(task_rows) - 1}")
for r in range(R0, R0 + len(task_rows)): ws.cell(r, 13).number_format = "0.0"

# ---------------- Code Sections ----------------
code_rows = []
for cs in D.CODE:
    cid, tid, module, layer, kind, name, sig, resp, feat, cin, cout, state = cs
    t = task_of[tid]
    code_rows.append((cid, tid, t[1], f"{t[2]} {cap_name[t[2]]}", module, layer, kind, name, sig, resp, feat, cin, cout, state,
                      ", ".join(tests_by_cs.get(cid, []))))
sheet("Code Sections", "Code sections", "Every class, function, route, graph node, MCP tool and test module the program defines.",
      ["ID", "Task", "Increment", "Capability", "Module", "Layer", "Kind", "Name", "Signature", "Responsibility",
       "Framework feature used", "Contracts in", "Contracts out", "Control state", "Tests"],
      [8, 7, 8, 22, 30, 12, 12, 26, 52, 44, 38, 30, 30, 18, 10], code_rows, code_cols=(9,))

# ---------------- Modules ----------------
mods = {}
for cs in D.CODE: mods.setdefault(cs[2], cs[3])
mod_rows = sorted(((m, l) for m, l in mods.items()), key=lambda x: x[0])
ws = sheet("Modules", "Modules", f"The {D.META['package']} package layout; section counts are live from the Code Sections sheet.",
           ["Module", "Layer", "Code sections", "Imports frameworks?"], [44, 14, 14, 40], [(m, l, None, None) for m, l in mod_rows], id_cols=(), code_cols=(1,))
rule = {"domain": "No — plain Python only", "contracts": "Pydantic only", "application": "No — ports are Protocols",
        "control": "No framework (Redis client only)", "agents": "LangGraph, AutoGen", "adapters": "Yes — one framework per adapter",
        "api": "FastAPI, Pydantic", "mcp_servers": "MCP (FastMCP)", "evals": "Eval frameworks", "tests": "pytest, Playwright",
        "config": "pydantic-settings", "workers": "Queue library", "observability": "OpenTelemetry"}
for r, (m, l) in enumerate(mod_rows, R0):
    ws.cell(r, 3, f"=COUNTIF('Code Sections'!$E${R0}:$E${LAST},A{r})").font = F_BODY
    ws.cell(r, 4, rule.get(l, "")).font = F_BODY
    for c in (3, 4): ws.cell(r, c).border = BORDER

# ---------------- Contracts ----------------
con_rows = []
for m in D.MODELS:
    con_rows.append((m[0], m[1], m[2], m[3], m[4], m[5], m[6], m[7], None, m[8],
                     ", ".join(sorted(set(cs_by_model.get(m[1], [])))), class_src(m)))
ws = sheet("Contracts", "Contracts (Pydantic models)", "Every model with a generated class skeleton. Field detail is on Contract Fields.",
           ["ID", "Model", "Module", "Kind", "Purpose", "Base", "model_config", "Validators / invariants", "Fields", "Defined in task", "Used by code sections", "Class skeleton"],
           [7, 22, 26, 20, 36, 18, 22, 34, 8, 10, 26, 90], con_rows, code_cols=(12,), freeze_col=3)
for r in range(R0, R0 + len(con_rows)):
    ws.cell(r, 9, f"=COUNTIF('Contract Fields'!$A${R0}:$A${LAST},B{r})").font = F_BODY
    ws.row_dimensions[r].height = min(400, 15 * (4 + len(D.FIELDS[ws.cell(r, 2).value])))

# ---------------- Contract Fields ----------------
mid = {m[1]: m[0] for m in D.MODELS}
f_rows = []
for m in D.MODELS:
    for i, (f, typ, default, cons, desc) in enumerate(D.FIELDS[m[1]], 1):
        f_rows.append((m[1], mid[m[1]], i, f, typ, default, cons, desc))
sheet("Contract Fields", "Contract fields", "One row per field. 'required' means no default.",
      ["Model", "Model ID", "#", "Field", "Type", "Default", "Constraints", "Meaning"],
      [22, 9, 5, 22, 48, 26, 26, 50], f_rows, id_cols=(1,), code_cols=(4, 5, 6))

# ---------------- remaining reference sheets ----------------
sheet("Ports & Adapters", "Ports and adapters", "Each framework sits behind one port; domain and application code never import it.",
      ["Port", "Methods", "Adapter", "Framework", "Features used", "Rule"], [16, 60, 20, 14, 60, 40], D.PORTS, code_cols=(2,))
sheet("Tool Registry", "Tool registry", "Every tool the run can call. Writes go only through control.act().",
      ["Tool", "Server", "Risk class", "Input model", "Output model", "Requires", "Idempotency key", "Verified by", "Undo", "Timeout (s)", "Vault scope"],
      [30, 9, 16, 24, 26, 40, 24, 20, 20, 10, 14], D.TOOLS, code_cols=(1,))
sheet("API Routes", "API routes", "FastAPI routes; the entity always comes from the token.",
      ["Method", "Path", "Request", "Response", "Dependencies", "Status codes", "Task"], [8, 34, 28, 28, 34, 30, 8], D.ROUTES, code_cols=(2,))
sheet("Run Graph", "Graphs: nodes and edges", "run_graph is the invoice run; resolution_graph and answer_graph are compiled subgraphs.",
      ["Graph", "Node", "Code section", "Control states", "Reads", "Writes", "Edges out"], [18, 14, 12, 26, 28, 26, 52], D.GRAPH)
sheet("Controls", "The 15 controls in code", "Where each control of the agent control model lives in this program.",
      ["Control", "What it means here", "Code sections", "Frameworks", "On failure"], [16, 60, 40, 26, 24], D.CONTROLS)
sheet("Recovery Policy", "Recovery policy", "One table owns every failure response. Limits are durable counters.",
      ["Error", "Raised by", "Response", "Limit", "Target", "After the limit"], [24, 28, 44, 7, 14, 32], D.RECOVERY)
sheet("Autonomy Matrix", "Autonomy matrix", "Thresholds come from Settings (post_alone_max_aed, two_approver_min_aed, daily_autonomous_cap_aed).",
      ["Action", "Risk class", "Runs alone when", "Needs approval when", "Never"], [34, 16, 44, 44, 30], D.AUTONOMY)
sheet("Persistence", "Persistence", "Who owns each table or store, its keys and its retention.",
      ["Table / store", "Technology", "Owner", "Key columns", "Purpose", "Retention"], [32, 20, 26, 60, 36, 24], D.PERSISTENCE, code_cols=(1, 4))
sheet("Events", "Events", "Payload models are on the Contracts sheet.",
      ["Event model", "Producer", "Consumers", "Transport"], [20, 34, 44, 14], D.EVENTS)
sheet("Tests", "Tests", "Every test names the code section it proves and the task that owns it.",
      ["ID", "Type", "Name", "Target code section", "Task", "Given / when / then"], [7, 11, 32, 12, 8, 90], D.TESTS)

# ---------------- Traceability ----------------
ws = wb.create_sheet("Traceability")
ws["A1"] = "Traceability"; ws["A1"].font = F_TITLE
ws["A2"] = "How many tasks build each control capability within each program capability. Live COUNTIFS over the Tasks sheet."; ws["A2"].font = F_SUB
caps = [c[0] for c in D.CAPABILITIES]
heads = ["Control capability"] + [f"{c}\n{cap_name[c]}" for c in caps] + ["Total"]
for c, h in enumerate(heads, 1):
    cell = ws.cell(HR, c, h); cell.font = F_HEAD; cell.fill = FILL_HEAD; cell.border = BORDER
    cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.column_dimensions[get_column_letter(c)].width = 28 if c == 1 else 13
ws.row_dimensions[HR].height = 44
for r, cc in enumerate(D.CONTROL_CAPS, R0):
    ws.cell(r, 1, cc).font = F_ID; ws.cell(r, 1).border = BORDER
    for c, cap in enumerate(caps, 2):
        cell = ws.cell(r, c, f'=COUNTIFS(Tasks!$C${R0}:$C${LAST},"{cap} *",Tasks!$H${R0}:$H${LAST},"*"&$A{r}&"*")')
        cell.font = F_BODY; cell.border = BORDER; cell.alignment = Alignment(horizontal="center")
    tc = len(caps) + 2
    cell = ws.cell(r, tc, f"=SUM({get_column_letter(2)}{r}:{get_column_letter(tc-1)}{r})")
    cell.font = Font(name="Arial", size=10, bold=True); cell.border = BORDER; cell.alignment = Alignment(horizontal="center")
last = R0 + len(D.CONTROL_CAPS) - 1
ws.conditional_formatting.add(f"B{R0}:{get_column_letter(len(caps)+1)}{last}",
                              ColorScaleRule(start_type="num", start_value=0, start_color="FFFFFF", end_type="max", end_color="F6C453"))
ws.freeze_panes = ws.cell(R0, 2)


# ======================= Claude Code sheets =======================
sheet("Stack", "Stack", "Minimum lines are what the code is written against; the resolved version (yellow) is copied from uv.lock at T-101.",
      ["Layer", "Component", "Package / install", "Minimum line", "Resolved version", "Purpose", "Used by tasks", "Docs", "Note"],
      [14, 22, 46, 16, 14, 40, 30, 34, 56], [r[:4] + ("",) + r[4:] for r in D.STACK], id_cols=(2,), edit_cols=(5,), freeze_col=3)
sheet("Framework Rules", "Framework rules", "Claude Code imports each framework only where this sheet allows. An import-linter contract enforces it.",
      ["Framework", "Import allowed only in", "Use it for", "Never use it for", "APIs to use", "You implement", "Gotchas"],
      [16, 30, 36, 36, 52, 40, 60], D.FRAMEWORK_RULES, code_cols=(2,))
sheet("Conventions", "Conventions", "Applies to every task.", ["ID", "Area", "Rule", "Example"], [7, 16, 80, 46], D.CONVENTIONS, code_cols=(4,))
sheet("Guardrails", "Guardrails", "Claude Code must never do these. They are repeated in CLAUDE.md.", ["ID", "Never", "Why"], [7, 90, 50], D.GUARDRAILS)
sheet("Definition of Done", "Definition of Done", "Every task, every time.", ["ID", "Check"], [7, 110], D.DOD)
sheet("Local Services", "Local services", "docker compose up -d starts the containers; the last four are processes run with uv.",
      ["Service", "Image / command", "Port", "Purpose", "First needed in"], [16, 52, 12, 44, 14], D.ENV_SERVICES, code_cols=(2,))
env_rows = [(f"ITP_{f.upper()}", typ, default, cons, desc) for f, typ, default, cons, desc in D.FIELDS["Settings"]]
sheet("Env Variables", "Environment variables", "Derived from the Settings contract (env_prefix ITP_). .env.example lists every one with a placeholder.",
      ["Variable", "Type", "Default", "Constraints", "Meaning"], [32, 20, 22, 24, 50], env_rows, code_cols=(1, 2, 3))
sheet("Commands", "Commands", "Make targets used by people, Claude Code and CI.", ["Target", "Command", "What it does"], [16, 100, 40], D.COMMANDS, code_cols=(2,))
sheet("Bootstrap Files", "Bootstrap files", "Created before feature work starts.", ["Path", "Created in", "Purpose", "Content"], [34, 10, 36, 100], D.BOOTSTRAP, code_cols=(1,))
sheet("Prompt Registry", "Prompt registry", "Prompts live in prompts/<id>/v<n>.md and are loaded by id; the version used is written to the ledger.",
      ["Prompt id", "Used by", "Model role (setting)", "Output schema", "Purpose", "Rules"], [22, 10, 20, 22, 46, 60], D.PROMPTS, code_cols=(1,))
sheet("Mail Templates", "Mail templates", "send_message only accepts these template ids under standing approval.", ["Template id", "Purpose", "Approval"], [26, 70, 22], D.MAIL_TEMPLATES, code_cols=(1,))
sheet("Fixtures", "Fixtures", "All values are fictitious test data. tests/fixtures/seed.py loads them into the stub ERP and the index.",
      ["Fixture", "Entity", "Values", "Used by tests"], [28, 18, 100, 16], D.FIXTURES)
sheet("Open Decisions", "Open decisions", "Claude Code uses the default until a decision is recorded (yellow).",
      ["ID", "Decision", "Why it matters", "Default until decided", "Owner", "Decision taken"], [8, 44, 36, 56, 10, 30],
      [r + ("",) for r in D.OPEN_DECISIONS], edit_cols=(6,))
sheet("Glossary", "Glossary", "", ["Term", "Meaning"], [24, 110], D.GLOSSARY)

# ---------------- Task prompts (generated) ----------------
fw_names = [r[0] for r in D.FRAMEWORK_RULES]
model_by_name = {m[1]: m for m in D.MODELS}
def task_prompt(t):
    tid, inc, cap, title, obj, fw, states, ccaps, dep, days, done = t
    secs = [c for c in D.CODE if c[1] == tid]
    used = sorted(contracts_by_task.get(tid, []), key=lambda m: model_order.get(m, 999))
    own = [m for m in used if model_by_name[m][8] == tid]
    other = [m for m in used if m not in own]
    ftext = (fw + ' ' + ' '.join(c[8] for c in secs)).lower()
    fws = [n for n in fw_names if n.lower().split()[0] in ftext]
    tests = [x for x in D.TESTS if x[4] == tid]
    L = [f"# {tid} · {title}", f"Increment {inc} · {cap} {cap_name[cap]} · estimate {days} day(s)", "",
         "Read CLAUDE.md first and follow its conventions and guardrails.",
         f"Depends on: {dep or 'nothing'} (must already be merged).", "", "## Goal", obj, "",
         "## Frameworks and features", fw,
         ("Framework rules that apply: " + ", ".join(fws)) if fws else "No framework: plain Python only.", "",
         "## Build these code sections"]
    for c in secs:
        L += [f"- {c[0]} `{c[2]}` · {c[4]} **{c[5]}**", f"  - signature: `{c[6]}`",
              f"  - responsibility: {c[7]}", f"  - framework feature: {c[8]}",
              f"  - contracts in: {c[9] or '—'} · out: {c[10] or '—'} · control state: {c[11]}"]
    if own:
        L += ["", "## Contracts defined in this task (implement exactly)"]
        for m in own: L += ["```python", class_src(model_by_name[m]), "```"]
    if other:
        L += ["", "## Contracts used (already defined; import, do not redefine)", ", ".join(other)]
    L += ["", "## Tests to write first (they must fail before your change)"]
    L += [f"- {x[0]} [{x[1]}] {x[2]} — {x[5]}" for x in tests] or ["- Add at least one unit test per code section; name them after the section."]
    L += ["", "## Controls delivered", f"Control states: {states} · control capabilities: {ccaps}", "",
          "## Done when", done, "", "## Out of scope",
          "Anything belonging to another task. Leave `TODO(T-xxx)` markers instead of building ahead.",
          "If something you need is not in the workbook, stop and ask instead of inventing it."]
    return "\n".join(L)
import build_repo_docs as RD
prompts = [(t[0], t[3], RD.task_prompt(t)) for t in D.TASKS]
ws = sheet("Task Prompts", "Task prompts for Claude Code", "Paste one into Claude Code, or use claude_code_pack/tasks/<id>.md. Generated; do not edit here.",
           ["Task", "Title", "Prompt"], [8, 30, 140], prompts, code_cols=(3,))
for r in range(R0, R0 + len(prompts)): ws.row_dimensions[r].height = 300

# ---------------- Repo docs + session protocol sheets ----------------
sheet("Repo Docs", "Repo docs", "Pushed with the code. Generated files are rebuilt by `make docs`; living files are created once and kept up to date by sessions.",
      ["Path", "Kind", "Read when", "Updated when", "Purpose"], [44, 11, 30, 36, 70], D.REPO_DOCS, code_cols=(1,))
PROTOCOL = [
    ("Start", "SessionStart hook", "Prints the 3-line progress summary (counts, current group, next group + model)", ".claude/settings.json · ~60 tokens"),
    ("Start", "CLAUDE.md + @docs/HANDOFF.md", "Auto-loaded: session rules, 11 condensed rules, the handoff", "~700 tokens"),
    ("Start", "/start-session", "progress.py summary + git status; read only the current task spec", "No whole-file reads of PROGRESS, PLAN, CONTRACTS, ARCHITECTURE"),
    ("Pick", "/next-group", "Next ready group; check the model; branch; progress.py start G-xx", "One branch per group; a group may take several sessions"),
    ("During", "progress.py done", "Marks the task done and appends a 3-line completion record", "The model never rewrites PROGRESS.md"),
    ("During", "/ask-question · /record-decision", "Append with cat >>; block the task if needed", "Never invent fields, tools or routes"),
    ("End", "/end-session", "done/block · DEBT rows · progress.py log · HANDOFF ≤200 words · progress.py check · commit", "Run even if unfinished"),
    ("Increment closed", "progress.py archive INC-n", "Moves done rows, records and log entries to docs/archive/INC-n.md", "Keeps living files small"),
    ("CI", "make docs-check", "progress.py check: statuses, dependencies, records, one group at a time", "Fails the build on drift"),
]
ws = sheet("Session Groups", "Session groups", "One Claude Code session (or a few) and one branch per group, in this order. Days and task counts are live from the Tasks sheet.",
           ["Group", "Name", "Tasks", "Model", "Why this model", "Tasks (#)", "Est. days"], [8, 28, 40, 9, 60, 10, 10],
           [g + (None, None) for g in D.SESSION_GROUPS])
for r in range(R0, R0 + len(D.SESSION_GROUPS)):
    ws.cell(r, 6, f"=SUMPRODUCT(--ISNUMBER(SEARCH(Tasks!$A${R0}:$A${R0 + len(D.TASKS) - 1},C{r})))").font = F_BODY
    c = ws.cell(r, 7, f"=SUMPRODUCT(--ISNUMBER(SEARCH(Tasks!$A${R0}:$A${R0 + len(D.TASKS) - 1},C{r})),Tasks!$M${R0}:$M${R0 + len(D.TASKS) - 1})")
    c.font = F_BODY; c.number_format = "0.0"
gt = R0 + len(D.SESSION_GROUPS)
ws.cell(gt, 2, "Total").font = Font(name="Arial", size=10, bold=True)
ws.cell(gt, 4, f'=COUNTIF(D{R0}:D{gt-1},"opus")&" opus · "&COUNTIF(D{R0}:D{gt-1},"sonnet")&" sonnet"').font = Font(name="Arial", size=10, bold=True)
for c, f in ((6, f"=SUM(F{R0}:F{gt-1})"), (7, f"=SUM(G{R0}:G{gt-1})")):
    ws.cell(gt, c, f).font = Font(name="Arial", size=10, bold=True)
sheet("Session Protocol", "Session protocol", "Nothing about the build lives in chat history: every session starts from and ends in the repo docs.",
      ["Phase", "Command / mechanism", "What it does", "Rule"], [12, 22, 90, 40], PROTOCOL, code_cols=(2,))

# ---------------- Repo starter: CLAUDE.md, docs/, tasks/, .claude/, scripts/, spec/ ----------------
import os, shutil, zipfile
pack = "/home/claude/itp/repo_starter"
shutil.rmtree(pack, ignore_errors=True)
RD.build(pack)
os.makedirs(f"{pack}/spec", exist_ok=True)
here = os.path.dirname(os.path.abspath(__file__))
for f in ("itp_data.py", "build_itp_workbook.py", "build_repo_docs.py", "progress_template.py"):
    shutil.copy(os.path.join(here, f), f"{pack}/spec/{f}")
zpath = os.environ.get("ITP_STARTER_ZIP", "/mnt/user-data/outputs/invoice_to_pay_repo_starter.zip")
with zipfile.ZipFile(zpath, "w", zipfile.ZIP_DEFLATED) as z:
    for dirpath, _, files in os.walk(pack):
        for f in sorted(files):
            full = os.path.join(dirpath, f)
            z.write(full, os.path.join("invoice_to_pay", os.path.relpath(full, pack)))

# order: Index, Claude Code group, then the spec sheets
order = ["Index"] + [s for s, _ in SHEETS]
wb._sheets = [wb[n] for n in order]

out = os.environ.get("ITP_WORKBOOK_OUT", "/mnt/user-data/outputs/InvoiceToPay_BuildWorkbook.xlsx")
wb.save(out); print("saved", out)
